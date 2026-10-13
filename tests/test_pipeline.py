"""Pipeline + diagnostics + SQL execution tests."""

from supply_chain_lab import diagnostics as D
from supply_chain_lab import features as F
from supply_chain_lab import pipeline as P


def test_headline_kpis_denominator(tables):
    agg = F.aggregate_items(tables["order_items"], tables["products"])
    ol = F.build_order_level(tables["orders"], agg, tables["customers"], tables["sellers"])
    k = D.headline_kpis(ol)
    assert k["n_eligible"] == int(ol["is_late"].notna().sum())
    assert abs(k["late_rate"] + k["on_time_rate"] - 1.0) < 1e-9


def test_sql_files_execute(tables):
    from supply_chain_lab.config import SQL_DIR

    for fname in ("data_quality.sql", "operational_kpis.sql", "seller_diagnostics.sql"):
        frames = D.run_sql_file(SQL_DIR / fname, tables)
        assert len(frames) > 0
        for f in frames:
            assert len(f) >= 0


def test_pipeline_demo_end_to_end(tmp_path):
    out = tmp_path / "outputs"
    res = P.run(
        mode="demo",
        out_dir=out,
        processed_dir=tmp_path / "processed",
        deck_path=tmp_path / "d.pptx",
    )
    assert res["mode"] == "demo"
    assert (out / "delivery_risk_report_demo.xlsx").exists()
    assert (out / "scenario_comparison.csv").exists()
    assert (out / "model_results.json").exists()


def test_pipeline_demo_hermetic(tmp_path):
    from supply_chain_lab.config import PROCESSED_DIR, REPORTS_DIR

    proc_before = (
        {p.name: p.stat().st_mtime_ns for p in PROCESSED_DIR.glob("*")}
        if PROCESSED_DIR.exists()
        else {}
    )
    deck_before = (
        {p.name: p.stat().st_mtime_ns for p in REPORTS_DIR.glob("*.pptx")}
        if REPORTS_DIR.exists()
        else {}
    )
    out = tmp_path / "outputs"
    proc_dir = tmp_path / "processed"
    deck_path = tmp_path / "deck.pptx"
    res = P.run(mode="demo", out_dir=out, processed_dir=proc_dir, deck_path=deck_path)
    assert res["mode"] == "demo"
    assert (proc_dir / "order_level_demo.csv").exists()
    assert deck_path.exists()
    proc_after = (
        {p.name: p.stat().st_mtime_ns for p in PROCESSED_DIR.glob("*")}
        if PROCESSED_DIR.exists()
        else {}
    )
    deck_after = (
        {p.name: p.stat().st_mtime_ns for p in REPORTS_DIR.glob("*.pptx")}
        if REPORTS_DIR.exists()
        else {}
    )
    assert proc_before == proc_after, "pipeline wrote outside tmp processed_dir"
    assert deck_before == deck_after, "pipeline wrote outside tmp deck_path"


def test_excel_values_chart_inputs(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "outputs"
    proc_dir = tmp_path / "processed"
    deck_path = tmp_path / "deck.pptx"
    res = P.run(mode="demo", out_dir=out, processed_dir=proc_dir, deck_path=deck_path)
    wb = load_workbook(out / "delivery_risk_report_demo.xlsx")
    ws = wb["Executive Dashboard"]
    # B6 = late_rate from pipeline KPIs
    assert abs(float(ws["B6"].value) - float(res["kpis"]["late_rate"])) < 1e-9
    ws_reg = wb["Regional Analysis"]
    assert len(ws_reg._charts) >= 1
    ws_sc = wb["Scenario Analysis"]
    assert float(ws_sc["B5"].value) == 2.0
    assert ws_sc["B5"].fill.fgColor.rgb.index("FFF2CC") >= 0 or "FFF2CC" in str(
        ws_sc["B5"].fill.fgColor.rgb
    )


def test_excel_workbook_sheets(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "outputs"
    P.run(
        mode="demo",
        out_dir=out,
        processed_dir=tmp_path / "processed",
        deck_path=tmp_path / "d.pptx",
    )
    wb = load_workbook(out / "delivery_risk_report_demo.xlsx")
    for sheet in (
        "Executive Dashboard",
        "Operational KPIs",
        "Seller Analysis",
        "Regional Analysis",
        "Delivery Delays",
        "Risk Prioritization",
        "Scenario Analysis",
        "Assumptions",
        "Data Dictionary",
    ):
        assert sheet in wb.sheetnames, f"missing sheet {sheet}"
