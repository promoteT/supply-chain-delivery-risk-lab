"""End-to-end pipeline: `python -m supply_chain_lab.pipeline --mode [auto|demo|real]`."""

from __future__ import annotations

import argparse
import json
import logging
from datetime import UTC, datetime
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

from . import diagnostics as D
from . import excel_reporting as XL
from . import features as F
from . import ingestion as I
from . import modelling as M
from . import simulation as S
from . import validation as V
from .config import (
    FIXTURES_DIR,
    OUTPUTS_DIR,
    PROCESSED_DIR,
    REPORTS_DIR,
    REVIEW_CAPACITIES,
    ScenarioAssumptions,
)

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
logger = logging.getLogger(__name__)


def ensure_fixtures() -> None:
    from .make_fixtures import write_fixtures

    if not (FIXTURES_DIR / "olist_orders_dataset.csv").exists():
        logger.info("Generating synthetic fixtures ...")
        write_fixtures()


def build_deck(
    out_path: Path,
    mode: str,
    kpis: dict,
    model_results: list[dict],
    scenario_df: pd.DataFrame,
    chart_paths: list[Path],
) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    demo_tag = " (DEMO — synthetic data)" if mode == "demo" else ""
    slides = [
        (
            "Business problem & approach",
            (
                f"Where do delays occur, what predicts them, and how to prioritize scarce review capacity?{demo_tag}\n"
                "Approach: quality-controlled pipeline → diagnostics → approval-time risk model → capacity simulation."
            ),
        ),
        (
            "Operational performance",
            (
                f"Eligible orders: {kpis.get('n_eligible')}; late rate: {kpis.get('late_rate')} | "
                f"mean lead time: {kpis.get('mean_lead_time_days')}d; "
                f"approval→carrier: {kpis.get('mean_approval_to_carrier_days')}d; "
                f"carrier→delivered: {kpis.get('mean_carrier_to_delivered_days')}d."
            ),
        ),
        (
            "Delay concentrations & potential drivers",
            (
                "Pareto + segment late rates show where delays concentrate. "
                "Treat as potential contributing factors — operational validation required."
            ),
        ),
        (
            "Early warning & capacity trade-offs",
            "Risk scores at approval time rank orders for review. "
            + (
                " | ".join(
                    f"{r.get('model')}: ROC-AUC {r.get('roc_auc')}, PR-AUC {r.get('pr_auc')}"
                    for r in model_results
                    if r.get("roc_auc")
                )
            )
            + ". "
            "Thresholds from validation; test held out.",
        ),
        (
            "Recommendations, limitations, next steps",
            "1) Prioritize top-risk reviews within daily capacity. 2) Validate seller/segment patterns operationally. "
            "3) Track rescued-rate before claiming savings. Limits: no inventory/cost/counterfactual data; "
            "undelivered orders excluded from late rates; fixture numbers are illustrative."
            if mode == "demo"
            else "1) Prioritize top-risk reviews within daily capacity. 2) Validate seller/segment patterns operationally. "
            "3) Track rescued-rate before claiming savings. Limits: no inventory/cost/counterfactual data; "
            "undelivered orders excluded from late rates.",
        ),
    ]
    for i, (title, body) in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        if slide.shapes.title is not None:
            tx = slide.shapes.title.text_frame
        else:
            tx = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.2), Inches(12), Inches(0.8)
            ).text_frame
        tx.text = f"{i + 1}. {title}"
        tx.paragraphs[0].runs[0].font.size = Pt(28)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(7.5), Inches(2)).text_frame
        box.word_wrap = True
        box.text = body
        if i in (1, 2) and chart_paths:
            pic = chart_paths[min(i - 1, len(chart_paths) - 1)]
            if pic.exists():
                slide.shapes.add_picture(str(pic), Inches(8.3), Inches(1.2), width=Inches(4.2))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path


def run(mode: str = "auto", out_dir: Path = OUTPUTS_DIR) -> dict:
    if mode == "demo":
        ensure_fixtures()
        tables, actual = I.load_tables(FIXTURES_DIR)
        actual = "demo"
    elif mode == "real":
        tables, actual = I.load_tables()
        if actual != "real":
            raise FileNotFoundError("Real Olist CSVs not found in data/raw/. See data/README.md.")
    else:
        try:
            tables, actual = I.load_tables()
        except FileNotFoundError:
            ensure_fixtures()
            tables, actual = I.load_tables(FIXTURES_DIR)
            actual = "demo"

    out_dir.mkdir(parents=True, exist_ok=True)
    charts_dir = out_dir / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)

    missing = I.check_mandatory_columns(tables)
    if missing:
        raise ValueError(f"Missing mandatory columns/tables: {missing}")

    orders = tables["orders"]
    for c in (
        "order_purchase_timestamp",
        "order_approved_at",
        "order_delivered_carrier_date",
        "order_delivered_customer_date",
        "order_estimated_delivery_date",
    ):
        if c in orders.columns:
            orders[c] = pd.to_datetime(orders[c], errors="coerce")

    val_report = V.validate_orders(orders)
    join_report = V.validate_join_cardinality(orders, tables["order_items"])
    _elig, _excluded, audit = V.split_eligible_orders(orders)
    miss = V.missingness_table(tables)

    items_agg = F.aggregate_items(tables["order_items"], tables.get("products"))
    order_level = F.build_order_level(
        orders, items_agg, tables.get("customers"), tables.get("sellers")
    )
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    order_level.to_csv(PROCESSED_DIR / f"order_level_{actual}.csv", index=False)

    period = (
        f"{pd.to_datetime(orders['order_purchase_timestamp']).min().date()} to "
        f"{pd.to_datetime(orders['order_purchase_timestamp']).max().date()}"
    )
    kpis = D.headline_kpis(order_level)
    seg_state = D.late_rate_by(order_level, "customer_state")
    seg_cat = D.late_rate_by(order_level, "primary_category")
    pareto = D.pareto_sellers(order_level, tables["order_items"])
    delay_table = (
        order_level.loc[order_level["is_late"] == True, ["delay_days"]]
        .describe(percentiles=[0.5, 0.75, 0.9, 0.95])
        .round(2)
        .reset_index()
        .rename(columns={"index": "statistic", "delay_days": "days_past_estimate"})
    )
    D.run_kpi_sql(  # SQL cross-check; raises if any query fails
        {k: v for k, v in tables.items() if k in ("orders", "customers", "order_items", "sellers")}
    )

    c1 = D.chart_lead_time_dist(order_level, charts_dir / "lead_time_dist.png", period, actual)
    D.chart_delay_severity(order_level, charts_dir / "delay_severity.png", period, actual)
    D.chart_late_by_segment(
        seg_state, "customer_state", charts_dir / "late_by_state.png", period, actual
    )
    c4 = D.chart_pareto(pareto, charts_dir / "pareto_sellers.png", period, actual)

    frame = F.build_model_frame(
        orders,
        tables["order_items"],
        tables.get("customers"),
        tables.get("sellers"),
        tables.get("products"),
    )
    M.assert_no_leakage(frame)
    splits = M.chronological_split(frame)
    models = M.build_models()
    model_results = [
        M.evaluate_model(n, m, splits, tuple(REVIEW_CAPACITIES)) for n, m in models.items()
    ]
    baselines = M.evaluate_rule_and_baseline(splits)
    test_scores = models["hgb"].predict_proba(splits["test"][M.FEATURE_COLS])[:, 1]
    rule_scores = M.business_rule_scores(splits["test"])
    comparison = S.compare_policies(splits["test"], test_scores, rule_scores)
    assumptions = ScenarioAssumptions()
    scenario_costed = S.hypothetical_costs(comparison, assumptions.cost_per_review)

    data_dict = pd.DataFrame(
        [
            {
                "kpi": "On-time delivery rate",
                "definition": "Eligible orders delivered on/before estimate / eligible orders",
                "denominator": "Eligible = delivered with actual + estimated dates",
                "unit": "share",
            },
            {
                "kpi": "Late-delivery rate",
                "definition": "Eligible orders delivered after estimate / eligible orders",
                "denominator": "Eligible only; missing/unreceived never counted as on-time",
                "unit": "share",
            },
            {
                "kpi": "Lead time",
                "definition": "order_delivered_customer_date - order_purchase_timestamp",
                "denominator": "Delivered orders with both timestamps",
                "unit": "days",
            },
            {
                "kpi": "Approval to carrier",
                "definition": "order_delivered_carrier_date - order_approved_at",
                "denominator": "Orders with both timestamps",
                "unit": "days",
            },
            {
                "kpi": "Carrier to delivered",
                "definition": "order_delivered_customer_date - order_delivered_carrier_date",
                "denominator": "Orders with both timestamps",
                "unit": "days",
            },
            {
                "kpi": "Delay severity",
                "definition": "order_delivered_customer_date - order_estimated_delivery_date (late only)",
                "denominator": "Late eligible orders",
                "unit": "days",
            },
            {
                "kpi": "Recall@k / Precision@k",
                "definition": "Among top-k risk-ranked test orders: share of all lates caught / share of reviewed that were late",
                "denominator": "Test set lates / k reviewed",
                "unit": "share",
            },
            {
                "kpi": "Expected rescued (scenario)",
                "definition": "HYPOTHETICAL: reviewed truly-late orders rescued with assumed effectiveness (Monte Carlo)",
                "denominator": "Illustrative; not observed",
                "unit": "orders",
            },
        ]
    )
    XL.build_workbook(
        out_dir / f"delivery_risk_report_{actual}.xlsx",
        actual,
        period,
        kpis,
        seg_state,
        seg_cat,
        pareto,
        model_results + baselines,
        scenario_costed,
        assumptions.describe(),
        data_dict,
        delay_table,
    )
    deck_path = build_deck(
        REPORTS_DIR / f"consulting_deck_{actual}.pptx",
        actual,
        kpis,
        model_results,
        scenario_costed,
        [c1, c4],
    )

    (out_dir / "validation_report.json").write_text(
        json.dumps(
            {
                "mode": actual,
                "run_at": datetime.now(UTC).isoformat(),
                "validation": val_report,
                "joins": join_report,
                "eligibility_audit": audit,
                "missing_columns": missing,
            },
            indent=2,
            default=str,
        )
    )
    (out_dir / "model_results.json").write_text(
        json.dumps(
            {"mode": actual, "models": model_results, "baselines": baselines}, indent=2, default=str
        )
    )
    comparison.to_csv(out_dir / "scenario_comparison.csv", index=False)
    seg_state.to_csv(out_dir / "late_by_state.csv", index=False)
    pareto.to_csv(out_dir / "pareto_sellers.csv", index=False)
    miss.to_csv(out_dir / "missingness.csv", index=False)

    logger.info("Pipeline complete (mode=%s). Artifacts in %s and %s", actual, out_dir, REPORTS_DIR)
    return {
        "mode": actual,
        "kpis": kpis,
        "model_results": model_results,
        "baselines": baselines,
        "deck": str(deck_path),
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--mode", choices=["auto", "demo", "real"], default="auto")
    ap.add_argument("--out", default=str(OUTPUTS_DIR))
    args = ap.parse_args()
    run(mode=args.mode, out_dir=Path(args.out))


if __name__ == "__main__":
    main()
